import json
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
import pandas as pd
import sys
from datetime import datetime

# --- Load Config ---
with open("config.json", "r") as f:
    config = json.load(f)

BRAND = config["branding"]
FORMATS = config["formats"]

# after you load `config = json.load(...)`
DEFAULT_BRANDING = {
    "primaryColor":   "#266B8E",
    "secondaryColor": "#FFC66C",
    "accentColor":    "#B7C4C0",
    "fontFamily":     "Calibri"
}

BRAND = config.get("branding", DEFAULT_BRANDING)

# pick whatever key is present:
primary_hex = BRAND.get("primary_color") or BRAND.get("primaryColor")
font_val    = BRAND.get("font")          or BRAND.get("fontFamily")

# strip the '#' if you need a pure hex:
PRIMARY_COLOR = primary_hex.lstrip("#")
FONT_NAME     = font_val
FORMATS = config.get("formats", {})
CURRENCY_FORMAT = FORMATS.get("currency", "$#,##0.00")
DATE_FORMAT     = FORMATS.get("date",     "MM/DD/YYYY")


# --- Predefined Styles ---
header_fill = PatternFill(start_color=PRIMARY_COLOR, end_color=PRIMARY_COLOR, fill_type="solid")
header_font = Font(color="FFFFFF", bold=True, name=FONT_NAME)
center_align = Alignment(horizontal="center", vertical="center")
thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)

# --- Load Data ---
census_df = pd.read_excel("uploaded_census.xlsx")
rate_df = pd.read_excel("uploaded_rates.xlsx", header=0)
renewal_df = pd.read_excel("uploaded_renewal_rates.xlsx", header=0)

# Clean headers
rate_df.columns = rate_df.columns.astype(str)
renewal_df.columns = renewal_df.columns.astype(str)
age_column = rate_df.columns[0]
plan_columns = rate_df.columns[1:]
renewal_plan_columns = renewal_df.columns[1:]

# Build current rate lookup
rate_lookup = {}
for plan in plan_columns:
    plan_rates = {}
    for _, row in rate_df.iterrows():
        try:
            age = int(row[age_column])
            rate = row[plan]
            plan_rates[age] = rate
        except (ValueError, TypeError):
            continue
    rate_lookup[plan] = plan_rates

# Build renewal rate lookup
renewal_lookup = {}
for plan in renewal_plan_columns:
    plan_rates = {}
    for _, row in renewal_df.iterrows():
        try:
            age = int(row[age_column])
            rate = row[plan]
            plan_rates[age] = rate
        except (ValueError, TypeError):
            continue
    renewal_lookup[plan] = plan_rates

# --- Calculate Age and Assign Rates ---
renewal_date = datetime(2025, 5, 1)
census_df['DOB'] = pd.to_datetime(census_df['DOB'], errors='coerce')
census_df['Age at Renewal'] = census_df['DOB'].apply(
    lambda dob: renewal_date.year - dob.year - ((renewal_date.month, renewal_date.day) < (dob.month, dob.day))
    if pd.notnull(dob) else None
)

# Ensure a Plan column exists
if 'Plan' not in census_df.columns:
    selected_plan = next(iter(rate_lookup))
    census_df['Plan'] = selected_plan

# Identify subscribers
tiers = []
subscriber_id = None
for _, row in census_df.iterrows():
    if row['Status'] == 'EE':
        subscriber_id = row.name
    tiers.append(subscriber_id)

census_df['SubscriberID'] = tiers

# Assign current and renewal rates per person
census_df['Current Rate'] = census_df.apply(
    lambda row: rate_lookup.get(row['Plan'], {}).get(row['Age at Renewal'], None), axis=1
)
census_df['Renewal Rate'] = census_df.apply(
    lambda row: renewal_lookup.get(row['Plan'], {}).get(row['Age at Renewal'], None), axis=1
)


# Contribution Options
contrib_type = sys.argv[1] if len(sys.argv) > 1 else "flat"
contrib_employee = float(sys.argv[2]) if len(sys.argv) > 2 else 0.0
contrib_dependent = float(sys.argv[3]) if len(sys.argv) > 3 else 0.0

CONTRIB_MODEL = {
    "type": contrib_type,
    "employee": contrib_employee,
    "dependent": contrib_dependent
}



# Group by Subscriber and calculate contributions
def apply_contribution(group):
    group = group.copy()
    total_rate = group['Current Rate'].sum()
    subscriber = group.iloc[0]
    n_dependents = len(group) - 1

    if CONTRIB_MODEL['type'] == 'flat':
        er_share = CONTRIB_MODEL['employee'] + (n_dependents * CONTRIB_MODEL['dependent'])
    elif CONTRIB_MODEL['type'] == 'percent':
        er_share = subscriber['Current Rate'] * (CONTRIB_MODEL['employee'] / 100.0)
        if n_dependents > 0:
            dep_rates = group.iloc[1:]['Current Rate'].sum()
            er_share += dep_rates * (CONTRIB_MODEL['dependent'] / 100.0)
    else:
        er_share = 0

    group.loc[:, 'Employer Share'] = er_share
    group.loc[:, 'Employee Share'] = total_rate - er_share
    return group

census_df = census_df.groupby('SubscriberID').apply(apply_contribution)
census_df.reset_index(drop=True, inplace=True)
census_df.index = range(len(census_df))

# Flag rows with no matches for debug output
census_df['Rate Match'] = census_df.apply(
    lambda row: "✅" if pd.notna(row['Current Rate']) and pd.notna(row['Renewal Rate']) else "❌", axis=1
)

# Clean and group
census_df['Current Rate'] = pd.to_numeric(census_df['Current Rate'], errors='coerce')
census_df['Renewal Rate'] = pd.to_numeric(census_df['Renewal Rate'], errors='coerce')
census_df['Employer Share'] = pd.to_numeric(census_df['Employer Share'], errors='coerce')
census_df['Employee Share'] = pd.to_numeric(census_df['Employee Share'], errors='coerce')

# --- Group by plan and age ---
grouped = census_df[census_df['Rate Match'] == '✅'].groupby(['Plan', 'Age at Renewal']).agg(
    Count=('Current Rate', 'size'),
    Current_Rate=('Current Rate', 'first'),
    Renewal_Rate=('Renewal Rate', 'first'),
    ER_Total=('Employer Share', 'sum'),
    EE_Total=('Employee Share', 'sum')
).reset_index()

grouped['Monthly Current'] = grouped['Count'] * grouped['Current_Rate']
grouped['Monthly Renewal'] = grouped['Count'] * grouped['Renewal_Rate']
grouped['$ Increase'] = grouped['Monthly Renewal'] - grouped['Monthly Current']
grouped['% Increase'] = grouped.apply(
    lambda row: ((row['Monthly Renewal'] - row['Monthly Current']) / row['Monthly Current']) if row['Monthly Current'] != 0 else None, axis=1
)

# --- Create Workbook ---
wb = Workbook()
ws = wb.active
ws.title = "Financial Summary"

headers = [
    "Plan Name", "Age Band", "Count", "Current Rate", "Renewal Rate",
    "Monthly Current", "Monthly Renewal", "$ Increase", "% Increase",
    "Employer Share", "Employee Share"
]

for col, text in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col, value=text)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = center_align
    cell.border = thin_border

for idx, row in grouped.iterrows():
    ws.cell(row=idx+2, column=1, value=row['Plan'])
    ws.cell(row=idx+2, column=2, value=row['Age at Renewal'])
    ws.cell(row=idx+2, column=3, value=row['Count'])
    ws.cell(row=idx+2, column=4, value=row['Current_Rate'])
    ws.cell(row=idx+2, column=5, value=row['Renewal_Rate'])
    ws.cell(row=idx+2, column=6, value=row['Monthly Current']).number_format = CURRENCY_FORMAT
    ws.cell(row=idx+2, column=7, value=row['Monthly Renewal']).number_format = CURRENCY_FORMAT
    ws.cell(row=idx+2, column=8, value=row['$ Increase']).number_format = CURRENCY_FORMAT
    ws.cell(row=idx+2, column=9, value=row['% Increase']).number_format = "0.00%"
    ws.cell(row=idx+2, column=10, value=row['ER_Total']).number_format = CURRENCY_FORMAT
    ws.cell(row=idx+2, column=11, value=row['EE_Total']).number_format = CURRENCY_FORMAT

# --- Total Row ---
total_row = len(grouped) + 2
ws.cell(row=total_row, column=1, value="Total")
ws.cell(row=total_row, column=3, value=grouped['Count'].sum())
ws.cell(row=total_row, column=6, value=grouped['Monthly Current'].sum()).number_format = CURRENCY_FORMAT
ws.cell(row=total_row, column=7, value=grouped['Monthly Renewal'].sum()).number_format = CURRENCY_FORMAT
ws.cell(row=total_row, column=8, value=grouped['$ Increase'].sum()).number_format = CURRENCY_FORMAT
ws.cell(row=total_row, column=9, value=(grouped['$ Increase'].sum() / grouped['Monthly Current'].sum())).number_format = "0.00%"
ws.cell(row=total_row, column=10, value=grouped['ER_Total'].sum()).number_format = CURRENCY_FORMAT
ws.cell(row=total_row, column=11, value=grouped['EE_Total'].sum()).number_format = CURRENCY_FORMAT

# --- Debug Tab ---
debug_ws = wb.create_sheet(title="Debug")
debug_headers = [
    "First Name", "Last Name", "DOB", "Age at Renewal", "Plan", "Status",
    "Current Rate", "Renewal Rate", "Rate Match", "Employer Share", "Employee Share"
]

for col, text in enumerate(debug_headers, start=1):
    cell = debug_ws.cell(row=1, column=col, value=text)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = center_align
    cell.border = thin_border

for idx, row in census_df.iterrows():
    debug_ws.cell(row=idx+2, column=1, value=row.get("First Name"))
    debug_ws.cell(row=idx+2, column=2, value=row.get("Last Name"))
    debug_ws.cell(row=idx+2, column=3, value=row.get("DOB"))
    debug_ws.cell(row=idx+2, column=4, value=row.get("Age at Renewal"))
    debug_ws.cell(row=idx+2, column=5, value=row.get("Plan"))
    debug_ws.cell(row=idx+2, column=6, value=row.get("Status"))
    debug_ws.cell(row=idx+2, column=7, value=row.get("Current Rate"))
    debug_ws.cell(row=idx+2, column=8, value=row.get("Renewal Rate"))
    debug_ws.cell(row=idx+2, column=9, value=row.get("Rate Match"))
    debug_ws.cell(row=idx+2, column=10, value=row.get("Employer Share"))
    debug_ws.cell(row=idx+2, column=11, value=row.get("Employee Share"))

from build_benefit_summary import build_benefit_summary

# Dummy plan data to test layout
plans = [
    {
        "name": "Current Plan",
        "benefits": {
            "Deductible (Individual)": "$500",
            "Deductible (Family)": "$1,000",
            "Out-of-Pocket Max (Individual)": "$3,000",
            "Out-of-Pocket Max (Family)": "$6,000",
            "PCP Visit": "$25",
            "Specialist Visit": "$40",
            "Urgent Care": "$50",
            "Emergency Room": "$300",
            "Generic Rx": "$10",
            "Brand Rx": "$35",
            "Specialty Rx": "$100",
            "Hospitalization": "$500/day",
            "Mental Health": "$25",
            "Telehealth": "$0"
        }
    },
    {
        "name": "Option 1",
        "benefits": {
            "Deductible (Individual)": "$750",
            "Deductible (Family)": "$1,500",
            "Out-of-Pocket Max (Individual)": "$4,000",
            "Out-of-Pocket Max (Family)": "$8,000",
            "PCP Visit": "$30",
            "Specialist Visit": "$45",
            "Urgent Care": "$60",
            "Emergency Room": "$350",
            "Generic Rx": "$15",
            "Brand Rx": "$40",
            "Specialty Rx": "$120",
            "Hospitalization": "$600/day",
            "Mental Health": "$30",
            "Telehealth": "$10"
        }
    }
]

# Create benefit summary worksheet
benefit_ws = wb.create_sheet(title="Benefit Summary")
build_benefit_summary(benefit_ws, plans)


# --- Save File ---
wb.save("financial_summary_output.xlsx")
